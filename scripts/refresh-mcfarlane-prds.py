"""Detect source/config drift and synchronize DOCX; never assert a review."""
from __future__ import annotations
import hashlib
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt
from monitor_backup import DOCS, manifest, fingerprint

LEDGER=DOCS/'.mcfarlane-prd-signatures.json'

def write_docx(markdown,dest):
    doc=Document(); doc.styles['Normal'].font.name='Aptos'; doc.styles['Normal'].font.size=Pt(10.5)
    for line in markdown.splitlines():
        if line.startswith('# '): doc.add_heading(line[2:],0)
        elif line.startswith('## '): doc.add_heading(line[3:],1)
        elif line.startswith('### '): doc.add_heading(line[4:],2)
        elif line.startswith('- '): doc.add_paragraph(line[2:],'List Bullet')
        elif line.strip(): doc.add_paragraph(re.sub(r'\*\*([^*]+)\*\*',r'\1',line))
    doc.core_properties.author=''; doc.core_properties.last_modified_by=''
    doc.save(dest)

def main():
    ledger=json.loads(LEDGER.read_text(encoding='utf-8')) if LEDGER.exists() else {}
    changed=[]
    for kind in ('exotic','sales'):
        stem=manifest()[kind]['prd']; md=DOCS/(stem+'.md'); dest=DOCS/(stem+'.docx')
        if not md.exists():
            continue  # A source-only restore may install only one monitor.
        sig=fingerprint(kind); text=md.read_text(encoding='utf-8')
        # Review baseline is deliberately not advanced by this automation.
        reviewed=re.search(r'\*\*Reviewed implementation fingerprint:\*\* `([0-9a-f]+)`',text)
        status='matches documented review baseline' if reviewed and reviewed[1]==sig else 'changed or unreviewed; substantive review required'
        block=f'<!-- automated-drift:start -->\n**Observed implementation fingerprint:** `{sig}`\n**Automated drift status:** {status}. Hash comparison is not a requirements review.\n<!-- automated-drift:end -->'
        text=re.sub(r'<!-- automated-drift:start -->.*?<!-- automated-drift:end -->',lambda m:block,text,flags=re.S) if '<!-- automated-drift:start -->' in text else text.rstrip()+'\n\n'+block+'\n'
        if text!=md.read_text(encoding='utf-8'): md.write_text(text,encoding='utf-8')
        mdhash=hashlib.sha256(text.encode()).hexdigest()
        prior=ledger.get(kind,{})
        if not isinstance(prior,dict) or prior.get('markdown')!=mdhash or not dest.exists():
            write_docx(text,dest); changed.append(kind)
        ledger[kind]={'observed':sig,'markdown':mdhash}
    serialized=json.dumps(ledger,indent=2)+'\n'
    if not LEDGER.exists() or LEDGER.read_text(encoding='utf-8')!=serialized: LEDGER.write_text(serialized,encoding='utf-8')
    print('PRD drift/DOCX synchronization: '+(', '.join(changed) if changed else 'no changes')+'; no automatic substantive review claimed')

if __name__=='__main__': main()
